"""Meeting-ready AI Action Plan export: PDF, PPTX, DOCX and XLSX.

Every format carries the same content, in this order: the two-sentence
summary, the recommendations, the reorder priorities table, four charts, and
a final line recording the filters and thresholds the report was built with.

Charts are drawn with matplotlib (Agg backend) and not exported from Plotly,
because Plotly's image export needs a headless Chrome (Kaleido) that is
fragile on Windows and in slim Docker images. The palette is imported from
pharmastock.charts, so the exports match the on-screen charts.

The PDF embeds DejaVu Sans, which ships inside matplotlib, so symbols such as
≤, ×, €, £ and ₹ render the same on Windows, macOS and Linux without
depending on system fonts.
"""

from __future__ import annotations

import io
import math
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from xml.sax.saxutils import escape

import matplotlib

matplotlib.use("Agg")  # headless: no GUI toolkit needed on any OS or in Docker

import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

from agents import ActionReport  # noqa: E402
from pharmastock import charts as ch  # noqa: E402
from pharmastock.analytics import EXPIRY_BUCKETS, STOCK_STATUSES, Analysis  # noqa: E402

TITLE = "PharmaStock AI Action Plan"
BRAND = "#1c5cab"
INK, INK_2, MUTED, LINE = "#0b0b0b", "#52514e", "#6f6d68", "#e1e0d9"
IMPACT_COLORS = {"High": ch.CRITICAL, "Medium": ch.WARNING, "Low": ch.BLUE[300]}
SUMMARY_HEADING = "Summary"
RECS_HEADING = "Recommendations"
REORDER_HEADING = "Reorder priorities"
REC_HEADER = ["#", "Action", "Impact", "When", "Why"]
DISCLAIMER = ("Decision support only. Clinical and regulatory decisions remain with the responsible pharmacist. "
              "AI-generated text can contain errors; verify figures against the source data.")


# ---------------------------------------------------------------------------
# Report content
# ---------------------------------------------------------------------------

@dataclass
class ReportBundle:
    dataset_label: str
    reference_date: str
    currency: str
    report: ActionReport
    charts: list[tuple[str, bytes]]
    reorder: pd.DataFrame
    scope_lines: list[str]
    generated_at: str = field(default_factory=lambda: datetime.now(timezone.utc).strftime("%d %b %Y %H:%M UTC"))

    @property
    def subtitle(self) -> str:
        return f"{self.dataset_label} · reference date {self.reference_date} · generated {self.generated_at}"


def _money(v: float, cur: str) -> str:
    if v is None or (isinstance(v, float) and (math.isnan(v) or math.isinf(v))):
        return "–"
    a = abs(v)
    if a >= 1_000_000:
        return f"{cur}{v / 1_000_000:,.2f}M"
    if a >= 100_000:
        return f"{cur}{v / 1_000:,.1f}K"
    return f"{cur}{v:,.0f}"


def _num(v: float, decimals: int = 1) -> str:
    if v is None or (isinstance(v, float) and (math.isnan(v) or math.isinf(v))):
        return "–"
    return f"{v:,.{decimals}f}"


def reorder_table(a: Analysis, cur: str, top: int = 15) -> pd.DataFrame:
    sku = a.sku
    pri = sku[sku["Priority"] != ""].sort_values(["Priority_Score", "Days_On_Hand"], ascending=[False, True]).head(top)
    return pd.DataFrame({
        "Priority": pri["Priority"].str.split(" ").str[0], "Drug": pri["Drug_Name"] + " " + pri["Strength"],
        "VEN": pri["VEN_Class"], "Status": pri["Stock_Status"],
        "Days on hand": pri["Days_On_Hand"].map(lambda v: _num(v, 1)),
        "Order qty": pri["Suggested_Order_Qty"].map(lambda v: _num(v, 0)),
        "Order value": pri["Suggested_Order_Value"].map(lambda v: _money(v, cur)), "Location": pri["Location"],
    }).reset_index(drop=True)


def build_bundle(a: Analysis, report: ActionReport, cur: str, dataset_label: str,
                 current_scope: str) -> ReportBundle:
    scope_lines = [report.scope or current_scope]
    if report.scope and report.scope != current_scope:
        scope_lines = [f"Recommendations: {report.scope}", f"Charts and reorder table: {current_scope}"]
    return ReportBundle(
        dataset_label=dataset_label, reference_date=f"{a.settings.reference_date:%d %b %Y}", currency=cur,
        report=report, charts=render_charts(a, cur), reorder=reorder_table(a, cur), scope_lines=scope_lines,
    )


def _rec_rows(r: ActionReport) -> list[list[str]]:
    return [[str(i), x["action"], x["impact"], x["timeframe"], x["rationale"]]
            for i, x in enumerate(r.recommendations, 1)]


# ---------------------------------------------------------------------------
# Charts (matplotlib -> PNG bytes)
# ---------------------------------------------------------------------------

def _style(ax, grid_axis: str = "x") -> None:
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color(ch.BASELINE)
    ax.tick_params(colors=MUTED, labelsize=9)
    ax.grid(axis=grid_axis, color=LINE, linewidth=0.8)
    ax.set_axisbelow(True)


def _png(fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return buf.getvalue()


def _empty() -> bytes:
    fig, ax = plt.subplots(figsize=(8, 3.6))
    ax.axis("off")
    ax.text(0.5, 0.5, "No data for the current filters", ha="center", va="center", color=MUTED, fontsize=12)
    return _png(fig)


def _hbar(labels, values, colors, text, xlabel: str) -> bytes:
    fig, ax = plt.subplots(figsize=(8, max(2.6, 0.42 * len(labels) + 0.9)))
    y = np.arange(len(labels))
    ax.barh(y, values, color=colors, height=0.62)
    ax.set_yticks(y, labels, color=INK_2, fontsize=9)
    ax.invert_yaxis()
    top = max(values) if len(values) and max(values) > 0 else 1
    ax.set_xlim(0, top * 1.2)
    for yi, v, t in zip(y, values, text):
        ax.text(v + top * 0.01, yi, t, va="center", fontsize=9, color=INK_2)
    ax.set_xlabel(xlabel, color=INK_2, fontsize=9)
    _style(ax)
    return _png(fig)


def chart_expiry(a: Analysis, cur: str) -> bytes:
    b = a.batches
    if b.empty:
        return _empty()
    g = b.groupby("Expiry_Bucket", observed=False)["Stock_Value"].sum().reindex(EXPIRY_BUCKETS).fillna(0)
    counts = b["Expiry_Bucket"].astype(str).value_counts()
    g = g[[(counts.get(k, 0) > 0) or k in EXPIRY_BUCKETS[:4] for k in g.index]]
    fig, ax = plt.subplots(figsize=(8, 3.8))
    x = np.arange(len(g))
    ax.bar(x, g.values, color=[ch.EXPIRY_COLORS[k] for k in g.index], width=0.62)
    ax.set_xticks(x, g.index, color=INK_2, fontsize=9)
    top = g.max() or 1
    ax.set_ylim(0, top * 1.18)
    for xi, v in zip(x, g.values):
        ax.text(xi, v + top * 0.02, _money(v, cur), ha="center", fontsize=9, color=INK_2)
    ax.set_ylabel(f"Stock value at cost ({cur.strip()})", color=INK_2, fontsize=9)
    _style(ax, "y")
    return _png(fig)


def chart_status(a: Analysis) -> bytes:
    g = a.sku["Stock_Status"].value_counts().reindex(STOCK_STATUSES).fillna(0)
    g = g[g > 0]
    if g.empty:
        return _empty()
    return _hbar(list(g.index), g.values, [ch.STATUS_COLORS[k] for k in g.index],
                 [f"{int(v)}" for v in g.values], "Number of SKUs")


def chart_category(a: Analysis, cur: str) -> bytes:
    g = a.sku.groupby("Therapeutic_Category")["Stock_Value"].sum().nlargest(10)
    if g.empty:
        return _empty()
    return _hbar(list(g.index), g.values, ch.BLUE[450], [_money(v, cur) for v in g.values],
                 f"Stock value at cost ({cur.strip()})")


def chart_abc_ven(a: Analysis) -> bytes:
    sku = a.sku
    if sku.empty:
        return _empty()
    m = pd.crosstab(sku["ABC_Class"], sku["VEN_Class"]).reindex(index=["A", "B", "C"], columns=["V", "E", "N"],
                                                                 fill_value=0)
    fig, ax = plt.subplots(figsize=(6, 3.8))
    ramp = matplotlib.colors.LinearSegmentedColormap.from_list("ps", [ch.BLUE[100], ch.BLUE[650]])
    ax.imshow(m.values, cmap=ramp, aspect="auto")
    ax.set_xticks(range(3), ["Vital", "Essential", "Non-essential"], color=INK_2, fontsize=9)
    ax.set_yticks(range(3), ["A (high value)", "B", "C (low value)"], color=INK_2, fontsize=9)
    hi = m.values.max() or 1
    for i in range(3):
        for j in range(3):
            v = int(m.values[i, j])
            ax.text(j, i, f"{v}\n{m.index[i]}{m.columns[j]}", ha="center", va="center", fontsize=10,
                    color="white" if v > hi * 0.55 else INK)
    for side in ax.spines.values():
        side.set_visible(False)
    ax.tick_params(length=0)
    return _png(fig)


def render_charts(a: Analysis, cur: str) -> list[tuple[str, bytes]]:
    return [
        ("Stock value by expiry window", chart_expiry(a, cur)),
        ("SKU stock status", chart_status(a)),
        ("Stock value by therapeutic category", chart_category(a, cur)),
        ("ABC-VEN matrix (SKU count)", chart_abc_ven(a)),
    ]


# ---------------------------------------------------------------------------
# PDF (reportlab)
# ---------------------------------------------------------------------------

def _register_pdf_fonts() -> tuple[str, str]:
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    if "DejaVuSans" not in pdfmetrics.getRegisteredFontNames():
        ttf = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
        try:
            pdfmetrics.registerFont(TTFont("DejaVuSans", str(ttf / "DejaVuSans.ttf")))
            pdfmetrics.registerFont(TTFont("DejaVuSans-Bold", str(ttf / "DejaVuSans-Bold.ttf")))
        except Exception:  # font file missing: fall back to built-in Helvetica
            return "Helvetica", "Helvetica-Bold"
        # lets <b> inside a Paragraph switch to the bold face
        addMapping("DejaVuSans", 0, 0, "DejaVuSans")
        addMapping("DejaVuSans", 1, 0, "DejaVuSans-Bold")
        addMapping("DejaVuSans", 0, 1, "DejaVuSans")
        addMapping("DejaVuSans", 1, 1, "DejaVuSans-Bold")
    return "DejaVuSans", "DejaVuSans-Bold"


def to_pdf(bundle: ReportBundle) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    regular, bold = _register_pdf_fonts()
    S = {
        "title": ParagraphStyle("t", fontName=bold, fontSize=20, leading=24, textColor=colors.HexColor(INK)),
        "meta": ParagraphStyle("m", fontName=regular, fontSize=8.5, leading=11, textColor=colors.HexColor(MUTED)),
        "h1": ParagraphStyle("h1", fontName=bold, fontSize=13.5, leading=17, spaceBefore=14, spaceAfter=6,
                             textColor=colors.HexColor(BRAND), keepWithNext=1),
        "summary": ParagraphStyle("s", fontName=regular, fontSize=11, leading=15.5, textColor=colors.HexColor(INK)),
        "cell": ParagraphStyle("c", fontName=regular, fontSize=8.5, leading=11, textColor=colors.HexColor(INK_2)),
        "cellb": ParagraphStyle("cb", fontName=bold, fontSize=8.5, leading=11, textColor=colors.HexColor(INK)),
        "scope": ParagraphStyle("sc", fontName=regular, fontSize=8, leading=10.5, textColor=colors.HexColor(INK_2)),
    }
    P = lambda text, style="cell": Paragraph(escape(str(text)), S[style])  # noqa: E731
    width = A4[0] - 3.6 * cm
    grid = colors.HexColor(LINE)

    def table(rows: list[list], col_widths: list[float]) -> Table:
        t = Table(rows, colWidths=col_widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, grid), ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef2f7"))]))
        return t

    r = bundle.report
    summary = Table([[P(r.summary, "summary")]], colWidths=[width])
    summary.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f1f6fd")),
                                 ("LINEBEFORE", (0, 0), (0, -1), 3, colors.HexColor(BRAND)),
                                 ("LEFTPADDING", (0, 0), (-1, -1), 10), ("TOPPADDING", (0, 0), (-1, -1), 8),
                                 ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    story: list = [P(TITLE, "title"), Spacer(1, 4), P(bundle.subtitle, "meta"),
                   P(SUMMARY_HEADING, "h1"), summary, P(RECS_HEADING, "h1")]

    rows = [[P(h, "cellb") for h in REC_HEADER]]
    rows += [[P(n), P(action, "cellb"), P(impact), P(when), P(why)] for n, action, impact, when, why in _rec_rows(r)]
    story.append(table(rows, [width * f for f in (0.05, 0.31, 0.10, 0.12, 0.42)]))

    story.append(P(REORDER_HEADING, "h1"))
    df = bundle.reorder
    if df.empty:
        story.append(P("No items need reordering."))
    else:
        rows = [[P(c, "cellb") for c in df.columns]] + [[P(v) for v in row] for row in df.astype(str).values]
        story.append(table(rows, [width * f for f in (0.08, 0.22, 0.06, 0.18, 0.1, 0.1, 0.11, 0.15)]))

    story.append(P("Charts", "h1"))
    for title, png in bundle.charts:
        img = Image(io.BytesIO(png))
        scale = min(width / img.imageWidth, 9.5 * cm / img.imageHeight)
        img.drawWidth, img.drawHeight = img.imageWidth * scale, img.imageHeight * scale
        story.append(KeepTogether([P(title, "cellb"), Spacer(1, 3), img, Spacer(1, 12)]))

    story += [Spacer(1, 6)] + [P(line, "scope") for line in bundle.scope_lines]
    story += [Spacer(1, 6), P(f"Generated by {r.model} · {r.created_at}. {DISCLAIMER}", "meta")]

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(regular, 7)
        canvas.setFillColor(colors.HexColor(MUTED))
        canvas.drawString(1.8 * cm, 1.1 * cm, f"{TITLE} · {bundle.dataset_label} · {bundle.reference_date}")
        canvas.drawRightString(A4[0] - 1.8 * cm, 1.1 * cm, f"Page {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm, topMargin=1.6 * cm,
                      bottomMargin=1.8 * cm, title=TITLE, author="PharmaStock").build(
        story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# DOCX (python-docx)
# ---------------------------------------------------------------------------

def to_docx(bundle: ReportBundle) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.8)
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(TITLE, level=0)
    meta = doc.add_paragraph().add_run(bundle.subtitle)
    meta.font.size = Pt(8.5)
    meta.font.color.rgb = RGBColor.from_string(MUTED[1:])

    def small_table(header: list[str], rows: list[list[str]], bold_col: int | None = None) -> None:
        t = doc.add_table(rows=1, cols=len(header))
        t.style = "Light Grid Accent 1"
        for cell, text in zip(t.rows[0].cells, header):
            cell.text = text
        for row in rows:
            for j, (cell, v) in enumerate(zip(t.add_row().cells, row)):
                cell.text = v
                for run in cell.paragraphs[0].runs:
                    run.font.size = Pt(9)
                    run.bold = j == bold_col

    r = bundle.report
    doc.add_heading(SUMMARY_HEADING, level=1)
    doc.add_paragraph().add_run(r.summary).font.size = Pt(11)
    doc.add_heading(RECS_HEADING, level=1)
    small_table(REC_HEADER, _rec_rows(r), bold_col=1)

    doc.add_heading(REORDER_HEADING, level=1)
    if bundle.reorder.empty:
        doc.add_paragraph("No items need reordering.")
    else:
        small_table(list(bundle.reorder.columns), bundle.reorder.astype(str).values.tolist())

    doc.add_heading("Charts", level=1)
    for title, png in bundle.charts:
        doc.add_paragraph().add_run(title).bold = True
        doc.add_picture(io.BytesIO(png), width=Inches(6.2))

    for line in bundle.scope_lines:
        doc.add_paragraph().add_run(line).font.size = Pt(8.5)
    note = doc.add_paragraph().add_run(f"Generated by {r.model} · {r.created_at}. {DISCLAIMER}")
    note.italic, note.font.size = True, Pt(8)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PPTX (python-pptx), 16:9
# ---------------------------------------------------------------------------

def to_pptx(bundle: ReportBundle) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu, Inches, Pt

    rgb = lambda h: RGBColor.from_string(h.lstrip("#"))  # noqa: E731
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def text(slide, x, y, w, h, runs: list[tuple[str, float, bool, str]], space_after: float = 6):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        for i, (t, size, bold, color) in enumerate(runs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(space_after)
            run = p.add_run()
            run.text = t
            run.font.size, run.font.bold, run.font.color.rgb = Pt(size), bold, rgb(color)
        return box

    def new_slide(title: str):
        s = prs.slides.add_slide(blank)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = rgb(BRAND)
        bar.line.fill.background()
        text(s, 0.6, 0.4, 12, 0.8, [(title, 28, True, INK)])
        text(s, 0.6, 7.0, 12.1, 0.35, [(f"PharmaStock · {bundle.dataset_label} · {bundle.reference_date}",
                                        9, False, MUTED)])
        return s

    def add_table(slide, top: float, header: list[str], rows: list[list[str]], widths: list[float],
                  size: int, bold_col: int | None = None, color_col: int | None = None) -> None:
        tbl = slide.shapes.add_table(len(rows) + 1, len(header), Inches(0.6), Inches(top), Inches(sum(widths)),
                                     Inches(0.4 * (len(rows) + 1))).table
        for j, (head, w) in enumerate(zip(header, widths)):
            tbl.columns[j].width = Emu(Inches(w))
            tbl.cell(0, j).text = head
        for i, row in enumerate(rows, 1):
            for j, v in enumerate(row):
                cell = tbl.cell(i, j)
                cell.text = v
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(size)
                        run.font.bold = j == bold_col
                        if j == color_col:
                            run.font.color.rgb = rgb(IMPACT_COLORS.get(v, INK_2))

    r = bundle.report

    # Title
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = rgb(BRAND)
    bg.line.fill.background()
    text(s, 0.9, 2.4, 11.5, 3, [(TITLE, 40, True, "#ffffff"), (bundle.subtitle, 14, False, "#cde2fb")],
         space_after=18)

    # Summary + recommendations: 4 on the first slide (under the summary), 6 on each continuation
    rows = _rec_rows(r)
    chunks = [rows[:4]] + [rows[i:i + 6] for i in range(4, len(rows), 6)]
    for n, chunk in enumerate(chunks):
        s = new_slide(RECS_HEADING if n == 0 else f"{RECS_HEADING} (continued)")
        top = 1.45
        if n == 0:
            text(s, 0.6, 1.3, 12.1, 0.9, [(r.summary, 15, False, INK_2)])
            top = 2.4
        add_table(s, top, REC_HEADER, chunk, [0.4, 4.0, 1.0, 1.3, 5.4], 11, bold_col=1, color_col=2)

    # Reorder priorities, 8 rows per slide
    df = bundle.reorder
    reorder_rows = df.astype(str).values.tolist()
    if not reorder_rows:
        s = new_slide(REORDER_HEADING)
        text(s, 0.6, 1.6, 12, 1, [("No items need reordering.", 16, False, INK_2)])
    for n in range(0, len(reorder_rows), 8):
        s = new_slide(REORDER_HEADING if n == 0 else f"{REORDER_HEADING} (continued)")
        add_table(s, 1.5, list(df.columns), reorder_rows[n:n + 8],
                  [1.0, 3.0, 0.7, 2.3, 1.2, 1.1, 1.3, 1.5], 11, bold_col=1)

    for title, png in bundle.charts:
        s = new_slide(title)
        pic = s.shapes.add_picture(io.BytesIO(png), Inches(0.6), Inches(1.45), height=Inches(5.3))
        if pic.width > Inches(12.1):  # keep inside the slide
            ratio = Inches(12.1) / pic.width
            pic.width, pic.height = Inches(12.1), int(pic.height * ratio)
        pic.left = int((prs.slide_width - pic.width) / 2)

    s = new_slide("Report scope")
    text(s, 0.6, 1.5, 12.1, 5.2, [(line, 14, False, INK_2) for line in bundle.scope_lines]
         + [(f"Generated by {r.model} · {r.created_at}. {DISCLAIMER}", 11, False, MUTED)], space_after=12)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# XLSX (openpyxl)
# ---------------------------------------------------------------------------

def to_xlsx(bundle: ReportBundle) -> bytes:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill

    wrap = Alignment(wrap_text=True, vertical="top")
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="EEF2F7")
    heading = Font(bold=True, size=12, color=BRAND[1:])
    footnote = Font(italic=True, size=9, color=MUTED[1:])
    r = bundle.report

    def header_row(ws, values: list[str]) -> None:
        ws.append(values)
        for c in ws[ws.max_row]:
            c.font, c.fill = bold, head_fill

    def scope_footer(ws) -> None:
        ws.append([])
        for line in bundle.scope_lines:
            ws.append([line])
            ws.cell(ws.max_row, 1).font = footnote

    wb = Workbook()
    ws = wb.active
    ws.title = "Action Plan"
    ws.append([TITLE])
    ws["A1"].font = Font(bold=True, size=16, color=BRAND[1:])
    ws.append([bundle.subtitle])
    ws.append([])
    ws.append([SUMMARY_HEADING])
    ws.cell(ws.max_row, 1).font = heading
    ws.append([r.summary])
    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=5)
    ws.cell(ws.max_row, 1).alignment = wrap
    ws.row_dimensions[ws.max_row].height = 48
    ws.append([])
    ws.append([RECS_HEADING])
    ws.cell(ws.max_row, 1).font = heading
    header_row(ws, REC_HEADER)
    for row in _rec_rows(r):
        ws.append([int(row[0])] + row[1:])
        ws.cell(ws.max_row, 2).font = bold
        for c in ws[ws.max_row]:
            c.alignment = wrap
    for col, w in zip("ABCDE", (5, 48, 10, 13, 70)):
        ws.column_dimensions[col].width = w
    scope_footer(ws)
    ws.append([f"Generated by {r.model} · {r.created_at}. {DISCLAIMER}"])
    ws.cell(ws.max_row, 1).font = footnote

    ws = wb.create_sheet(REORDER_HEADING)
    header_row(ws, list(bundle.reorder.columns))
    for row in bundle.reorder.astype(str).values.tolist():
        ws.append(row)
    for col, w in zip("ABCDEFGH", (10, 34, 6, 26, 13, 11, 13, 22)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    scope_footer(ws)

    ws = wb.create_sheet("Charts")
    row = 1
    for title, png in bundle.charts:
        ws.cell(row, 1, title).font = Font(bold=True, size=12)
        img = XLImage(io.BytesIO(png))
        scale = 620 / img.width
        img.width, img.height = int(img.width * scale), int(img.height * scale)
        ws.add_image(img, f"A{row + 1}")
        row += int(img.height / 20) + 4
    for line in bundle.scope_lines:
        row += 1
        ws.cell(row, 1, line).font = footnote

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


EXPORTERS = {
    "PDF": (to_pdf, "pdf", "application/pdf"),
    "PPTX": (to_pptx, "pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
    "DOCX": (to_docx, "docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "XLSX": (to_xlsx, "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
}


def export_all(bundle: ReportBundle) -> dict[str, bytes]:
    return {fmt: fn(bundle) for fmt, (fn, _, _) in EXPORTERS.items()}
